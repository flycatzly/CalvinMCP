#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
read_cases.py — 把 Excel/CSV 手工用例表读成 JSON

为什么用 Python 而不是给 Node 装一个 xlsx 依赖：
  这个项目的 MCP 是零依赖设计（只为协议手写 stdio 循环，只装必要的驱动）。
  而 DSH 运行时自带 Python 与 openpyxl，Excel 读取交给它比往 Node 里塞一个
  上百个传递依赖的表格库更划算。openpyxl 缺失时自动回落到内置 zip+XML 解析
  （标准库，零第三方依赖）；两者都不可用时，本工具会明确报出缺什么。

用法：
    python read_cases.py <输入文件> --out <输出.json> [--sheet <表名>]

输出 JSON 形状：
    {
      "file": "...", "sheets": [...], "sheet": "Sheet1",
      "headerRow": 1, "columns": {...},
      "cases": [ { "row": 2, "title": "...", "steps": [...], "expect": "...",
                   "env": "test", "account": "u1", "raw": {...} } ],
      "warnings": [...]
    }
"""
import argparse
import csv
import json
import os
import sys

# 表头别名：中英文都认，团队的手工用例表两种写法都常见
ALIASES = {
    "title": ["用例", "用例名称", "用例标题", "场景", "场景名称", "case", "title", "name", "test case"],
    "steps": ["步骤", "操作步骤", "测试步骤", "执行步骤", "steps", "step", "操作"],
    "expect": ["预期", "预期结果", "期望结果", "验收点", "断言", "expect", "expected", "expected result"],
    "env": ["环境", "env", "environment"],
    "account": ["账号", "账户", "用户", "account", "user"],
    "priority": ["优先级", "级别", "priority"],
    "id": ["编号", "用例编号", "id", "no", "序号"],
    "precondition": ["前置条件", "前提", "precondition", "given"],
}


def norm(s):
    return str(s or "").strip().lower()


def match_column(header_cell):
    h = norm(header_cell)
    if not h:
        return None
    for key, names in ALIASES.items():
        for n in names:
            if h == norm(n) or norm(n) in h:
                return key
    return None


def split_steps(raw):
    """步骤单元格 -> 步骤列表。支持换行、编号、分号分隔。"""
    if raw is None:
        return []
    text = str(raw).replace("\r\n", "\n").replace("\r", "\n")
    # 先按换行切；若只有一行，再按分号/编号切
    parts = [p for p in text.split("\n") if p.strip()]
    if len(parts) <= 1:
        import re
        parts = re.split(r"\s*(?:;|；|->|→|\d+[\.、)）])\s*", text)
    cleaned = []
    for p in parts:
        s = str(p).strip()
        # 去掉行首的 1. / 1、 / (1) / - 等编号
        import re
        s = re.sub(r"^\s*(?:[-*•]|\(?\d+[\.、)）])\s*", "", s).strip()
        if s:
            cleaned.append(s)
    return cleaned


def find_header(rows, max_scan=10):
    """在前若干行里找表头行：命中别名最多的那一行。"""
    best_idx, best_score, best_map = -1, 0, {}
    for i, row in enumerate(rows[:max_scan]):
        colmap = {}
        for j, cell in enumerate(row):
            k = match_column(cell)
            if k and k not in colmap:
                colmap[k] = j
        score = len(colmap)
        if score > best_score:
            best_idx, best_score, best_map = i, score, colmap
    # 至少要认出 2 列，否则不认为这是用例表
    if best_score < 2:
        return -1, {}
    return best_idx, best_map


def read_xlsx_stdlib(path, sheet=None):
    """openpyxl 不可用时的兜底：xlsx 本质是 zip + XML，标准库足够解析。

    为什么值得内置（实测踩到的坑）：CI 容器和新机器的 Python 都没有 openpyxl，
    于是 Excel 编排链路在那里**整条挂掉**——而挂掉的方式像「环境没配好」，
    很容易被归到「回头再说」，其实只需要一个 zip 解析。

    覆盖用例表场景足够：共享字符串、行内字符串、布尔、数值、公式缓存值。
    不覆盖：日期序列号转日期（用例表几乎不用；openpyxl 装上时走原路径不受影响）。
    """
    import re
    import zipfile
    import xml.etree.ElementTree as ET

    NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    RNS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

    with zipfile.ZipFile(path) as z:
        shared = []
        if "xl/sharedStrings.xml" in z.namelist():
            root = ET.fromstring(z.read("xl/sharedStrings.xml"))
            for si in root.findall(NS + "si"):
                shared.append("".join(t.text or "" for t in si.iter(NS + "t")))

        wb = ET.fromstring(z.read("xl/workbook.xml"))
        sheets_meta = [(sh.get("name"), sh.get(RNS + "id")) for sh in wb.find(NS + "sheets")]
        if not sheets_meta:
            return None, "xlsx 里没有任何工作表"

        rels = {}
        for rel in ET.fromstring(z.read("xl/_rels/workbook.xml.rels")):
            rels[rel.get("Id")] = rel.get("Target")

        def sheet_target(rid):
            t = rels.get(rid) or ""
            if t.startswith("/"):
                return t.lstrip("/")
            return t if t.startswith("xl/") else "xl/" + t

        names = [n for n, _ in sheets_meta]
        pick = next(((n, rid) for n, rid in sheets_meta if sheet and n == sheet), sheets_meta[0])
        rows_root = ET.fromstring(z.read(sheet_target(pick[1])))

        def col_index(ref):
            m = re.match(r"([A-Za-z]+)", ref or "")
            n = 0
            for ch in (m.group(1) if m else ""):
                n = n * 26 + (ord(ch.upper()) - 64)
            return n - 1

        def cell_value(c):
            t = c.get("t")
            if t == "inlineStr":
                is_el = c.find(NS + "is")
                return "".join(x.text or "" for x in is_el.iter(NS + "t")) if is_el is not None else ""
            v_el = c.find(NS + "v")
            v = v_el.text if v_el is not None else None
            if v is None:
                return None
            if t == "s":
                try:
                    return shared[int(v)]
                except (ValueError, IndexError):
                    return v
            if t == "b":
                return v == "1"
            if t in ("str", "e"):
                return v
            try:
                f = float(v)
                return int(f) if f.is_integer() else f
            except ValueError:
                return v

        dense = []
        max_col = 0
        for row_el in rows_root.iter(NS + "row"):
            ri = int(row_el.get("r") or len(dense) + 1) - 1
            while len(dense) < ri:
                dense.append([])
            cells = {}
            for c in row_el.findall(NS + "c"):
                ci = col_index(c.get("r"))
                cells[ci] = cell_value(c)
                max_col = max(max_col, ci + 1)
            dense.append([cells.get(i) for i in range(max(cells) + 1)] if cells else [])

        # 补齐列宽：openpyxl 的 values_only 每行等长，兜底必须形状一致，
        # 否则下游「按列索引取值」会因为行短一截取到 None。
        rows = [row + [None] * (max_col - len(row)) for row in dense]
        return {"sheets": names, "sheet": pick[0], "rows": rows}, None


def read_xlsx(path, sheet=None):
    try:
        from openpyxl import load_workbook
    except ImportError:
        # 兜底：缺可选依赖不该让 Excel 链路整条失效（零依赖承诺的一部分）
        try:
            return read_xlsx_stdlib(path, sheet)
        except Exception as e:
            return None, "缺少 openpyxl，内置解析也失败了：%s" % e
    wb = load_workbook(path, data_only=True, read_only=True)
    names = list(wb.sheetnames)
    ws = wb[sheet] if sheet and sheet in names else wb[names[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return {"sheets": names, "sheet": ws.title, "rows": rows}, None


def read_csv_file(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except Exception:
            dialect = csv.excel
        rows = [list(r) for r in csv.reader(f, dialect)]
    return {"sheets": ["(csv)"], "sheet": "(csv)", "rows": rows}, None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("--out", required=True)
    ap.add_argument("--sheet", default=None)
    args = ap.parse_args()

    path = args.input
    if not os.path.exists(path):
        json.dump({"error": "文件不存在: %s" % path}, open(args.out, "w", encoding="utf-8"), ensure_ascii=False)
        return 2

    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        data, err = read_xlsx(path, args.sheet)
    elif ext == ".csv":
        data, err = read_csv_file(path)
    else:
        data, err = None, "不支持的扩展名 %s（支持 .xlsx / .xlsm / .csv；.xls 请先另存为 .xlsx）" % ext

    if err:
        json.dump({"error": err}, open(args.out, "w", encoding="utf-8"), ensure_ascii=False)
        return 2

    rows = data["rows"]
    warnings = []
    header_idx, colmap = find_header(rows)
    if header_idx < 0:
        json.dump({
            "file": path, "sheets": data["sheets"], "sheet": data["sheet"],
            "error": "没能识别出表头行。需要至少两列能被识别：用例/场景、步骤、预期结果、环境、账号。",
            "warnings": warnings, "cases": [],
        }, open(args.out, "w", encoding="utf-8"), ensure_ascii=False)
        return 3

    cases = []
    for ri in range(header_idx + 1, len(rows)):
        row = rows[ri]
        if not any(str(c).strip() for c in row if c is not None):
            continue
        get = lambda key: (row[colmap[key]] if key in colmap and colmap[key] < len(row) else None)
        title = str(get("title") or "").strip()
        steps = split_steps(get("steps"))
        expect = get("expect")
        if not title and not steps:
            continue
        if not title:
            title = "第 %d 行（未命名）" % (ri + 1)
            warnings.append("第 %d 行没有用例名称，已用行号占位。" % (ri + 1))
        if not steps:
            warnings.append("用例「%s」没有步骤，无法生成脚本。" % title)
        if not expect or not str(expect).strip():
            warnings.append("用例「%s」没有预期结果 —— 没有断言的用例只证明「页面没崩」，不证明任何业务事实。" % title)
        cases.append({
            "row": ri + 1,
            "id": str(get("id") or "").strip(),
            "title": title,
            "steps": steps,
            "expect": str(expect).strip() if expect is not None else "",
            "env": str(get("env") or "").strip(),
            "account": str(get("account") or "").strip(),
            "priority": str(get("priority") or "").strip(),
            "precondition": str(get("precondition") or "").strip(),
            "raw": {k: (str(v) if v is not None else "") for k, v in zip(
                ["c%d" % i for i in range(len(row))], row)},
        })

    out = {
        "file": path,
        "sheets": data["sheets"],
        "sheet": data["sheet"],
        "headerRow": header_idx + 1,
        "columns": {k: (v + 1) for k, v in colmap.items()},
        "caseCount": len(cases),
        "cases": cases,
        "warnings": warnings,
    }
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    return 0


if __name__ == "__main__":
    sys.exit(main())
